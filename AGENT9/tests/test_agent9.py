# =============================================================================
# tests/test_agent9.py
# Tests pytest pour Agent 9 — couverture minimale 60% sur modules critiques
#
# Pour lancer : pytest tests/test_agent9.py -v
# Pour voir la couverture : pytest tests/test_agent9.py --cov=. --cov-report=term
# =============================================================================

import os
import sys
import json
import zipfile
import tempfile
import pytest

# Ajouter le dossier parent au PYTHONPATH pour importer nos modules
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from fake_data import FAKE_PIPELINE_DATA
from models import PipelineData, ProductSpecs
from generators.html_gen import generate_html
from generators.json_gen import generate_json
from generators.excel_gen import generate_excel
from generators.xml_gen import generate_xml
from agent9 import run_agent9, _sha256, _create_manifest


# =============================================================================
# Fixtures — données réutilisées dans tous les tests
# =============================================================================

@pytest.fixture
def valid_data() -> PipelineData:
    """Retourne un PipelineData valide depuis les fake data."""
    return PipelineData(**FAKE_PIPELINE_DATA)


@pytest.fixture
def jinja_context(valid_data) -> dict:
    """Retourne le contexte Jinja2 utilisé par les templates HTML/XML."""
    return {
        "specs":          valid_data.agent1,
        "cad":            valid_data.agent2,
        "video":          valid_data.agent3,
        "sourcing":       valid_data.agent4,
        "nego":           valid_data.agent5,
        "tco":            valid_data.agent6,
        "bp":             valid_data.agent7,
        "twin":           valid_data.agent8,
        "generated_date": "2026-04-12 09:00:00",
    }


@pytest.fixture
def tmp_dir():
    """Fournit un dossier temporaire nettoyé après chaque test."""
    with tempfile.TemporaryDirectory() as d:
        yield d


# =============================================================================
# Tests des modèles Pydantic (validation données entrée)
# =============================================================================

class TestModels:

    def test_pipeline_data_valid(self):
        """Le PipelineData complet se construit sans erreur."""
        data = PipelineData(**FAKE_PIPELINE_DATA)
        assert data.agent1.part_id == "VLV-DN100-PN40-316L"

    def test_product_specs_fields(self, valid_data):
        """Les champs numériques sont correctement typés."""
        specs = valid_data.agent1
        assert isinstance(specs.diameter_mm, float)
        assert isinstance(specs.pressure_bar, float)
        assert specs.diameter_mm == 100.0
        assert specs.pressure_bar == 40.0

    def test_reliability_score_range(self, valid_data):
        """Le score de fiabilité doit être entre 0 et 100."""
        for s in valid_data.agent4.top_suppliers:
            assert 0 <= s.reliability_score <= 100

    def test_missing_required_field_raises(self):
        """Un champ obligatoire manquant doit lever une ValidationError."""
        bad_data = FAKE_PIPELINE_DATA.copy()
        bad_data["agent1"] = {"part_name": "Test"}  # champs manquants
        with pytest.raises(Exception):  # ValidationError de Pydantic
            PipelineData(**bad_data)

    def test_swot_has_all_categories(self, valid_data):
        """Le SWOT doit avoir les 4 catégories non vides."""
        swot = valid_data.agent7.swot
        assert len(swot.strengths) > 0
        assert len(swot.weaknesses) > 0
        assert len(swot.opportunities) > 0
        assert len(swot.threats) > 0

    def test_twin_alerts_severity_values(self, valid_data):
        """Les alertes ne peuvent avoir que 'critical' ou 'warning'."""
        valid_severities = {"critical", "warning"}
        for alert in valid_data.agent8.alerts:
            assert alert.severity in valid_severities


# =============================================================================
# Tests du générateur HTML
# =============================================================================

class TestHTMLGenerator:

    def test_html_file_created(self, jinja_context, tmp_dir):
        """Le fichier HTML est bien créé."""
        path = os.path.join(tmp_dir, "test.html")
        result = generate_html(jinja_context, path)
        assert os.path.exists(result)

    def test_html_not_empty(self, jinja_context, tmp_dir):
        """Le fichier HTML n'est pas vide."""
        path = os.path.join(tmp_dir, "test.html")
        generate_html(jinja_context, path)
        assert os.path.getsize(path) > 1000  # Au moins 1 KB

    def test_html_contains_product_name(self, jinja_context, tmp_dir):
        """Le HTML contient le nom du produit."""
        path = os.path.join(tmp_dir, "test.html")
        generate_html(jinja_context, path)
        content = open(path, encoding="utf-8").read()
        assert "Vanne Industrielle DN100 PN40" in content

    def test_html_contains_supplier(self, jinja_context, tmp_dir):
        """Le HTML contient la mention du réseau de fabrication."""
        path = os.path.join(tmp_dir, "test.html")
        generate_html(jinja_context, path)
        content = open(path, encoding="utf-8").read()
        assert "certified industrial manufacturing network" in content

    def test_html_is_valid_utf8(self, jinja_context, tmp_dir):
        """Le HTML est encodé en UTF-8 valide (accents, caractères spéciaux)."""
        path = os.path.join(tmp_dir, "test.html")
        generate_html(jinja_context, path)
        # Ne doit pas lever d'exception
        content = open(path, encoding="utf-8").read()
        assert "é" in content or "è" in content  # accents présents


# =============================================================================
# Tests du générateur JSON
# =============================================================================

class TestJSONGenerator:

    def test_json_file_created(self, valid_data, tmp_dir):
        """Le fichier JSON est bien créé."""
        path = os.path.join(tmp_dir, "test.json")
        result = generate_json(valid_data, path)
        assert os.path.exists(result)

    def test_json_is_valid(self, valid_data, tmp_dir):
        """Le JSON produit est syntaxiquement valide."""
        path = os.path.join(tmp_dir, "test.json")
        generate_json(valid_data, path)
        with open(path, encoding="utf-8") as f:
            data = json.load(f)  # Lève json.JSONDecodeError si invalide
        assert isinstance(data, dict)

    def test_json_is_strictly_commercial(self, valid_data, tmp_dir):
        """Le JSON ne contient que la partie commerciale et aucune donnée technique interne."""
        path = os.path.join(tmp_dir, "test.json")
        generate_json(valid_data, path)
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        
        # Vérifier la présence de la clé principale
        assert "commercial_catalogue" in data
        assert "catalogue_version" in data
        
        # Vérifier l'ABSENCE des clés internes
        assert "summary" not in data
        assert "full_data" not in data
        
        # Vérifier que les noms de fournisseurs ne leakent pas
        json_str = json.dumps(data)
        assert "AcierPro SARL" not in json_str

    def test_json_structure_compliance(self, valid_data, tmp_dir):
        """Le JSON respecte la structure des 10 sections commerciales (version détaillée)."""
        path = os.path.join(tmp_dir, "test.json")
        generate_json(valid_data, path)
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        
        sections = data["commercial_catalogue"]
        assert "1_executive_summary" in sections
        assert "3_detailed_specifications" in sections
        assert "5_pricing" in sections
        
        # Vérifier que les points sont bien des listes
        assert isinstance(sections["1_executive_summary"]["commercial_highlights"], list)
        assert isinstance(sections["9_client_value_justification"]["benefits"], list)

    def test_json_tco_value_correct(self, valid_data, tmp_dir):
        """La valeur TCO dans le JSON correspond à celle des données."""
        path = os.path.join(tmp_dir, "test.json")
        generate_json(valid_data, path)
        with open(path, encoding="utf-8") as f:
            data = json.load(f)
        # La valeur est maintenant dans Pricing
        assert data["commercial_catalogue"]["5_pricing"]["final_total_price_usd"] == 259200.0


# =============================================================================
# Tests du générateur Excel
# =============================================================================

class TestExcelGenerator:

    def test_excel_file_created(self, valid_data, tmp_dir):
        """Le fichier Excel est bien créé."""
        path = os.path.join(tmp_dir, "test.xlsx")
        result = generate_excel(valid_data, path)
        assert os.path.exists(result)

    def test_excel_not_empty(self, valid_data, tmp_dir):
        """Le fichier Excel n'est pas vide (> 5 KB)."""
        path = os.path.join(tmp_dir, "test.xlsx")
        generate_excel(valid_data, path)
        assert os.path.getsize(path) > 5000

    def test_excel_has_6_sheets(self, valid_data, tmp_dir):
        """Le fichier Excel a exactement 6 onglets."""
        from openpyxl import load_workbook
        path = os.path.join(tmp_dir, "test.xlsx")
        generate_excel(valid_data, path)
        wb = load_workbook(path)
        assert len(wb.sheetnames) == 6

    def test_excel_sheet_names(self, valid_data, tmp_dir):
        """Les onglets ont les bons noms."""
        from openpyxl import load_workbook
        path = os.path.join(tmp_dir, "test.xlsx")
        generate_excel(valid_data, path)
        wb = load_workbook(path)
        names = wb.sheetnames
        # Vérifier que les onglets essentiels sont présents
        assert any("Résumé" in n or "Resume" in n for n in names)
        assert any("TCO" in n for n in names)
        assert any("Business" in n for n in names)


# =============================================================================
# Tests du générateur XML
# =============================================================================

class TestXMLGenerator:

    def test_xml_file_created(self, jinja_context, tmp_dir):
        """Le fichier XML est bien créé."""
        path = os.path.join(tmp_dir, "test.xml")
        result = generate_xml(jinja_context, path)
        assert os.path.exists(result)

    def test_xml_is_valid(self, jinja_context, tmp_dir):
        """Le XML produit est syntaxiquement valide (lxml)."""
        from lxml import etree
        path = os.path.join(tmp_dir, "test.xml")
        generate_xml(jinja_context, path)
        # Doit parser sans exception
        tree = etree.parse(path)
        assert tree.getroot() is not None

    def test_xml_root_tag(self, jinja_context, tmp_dir):
        """Le tag racine du XML est 'IndustrialCatalogue'."""
        from lxml import etree
        path = os.path.join(tmp_dir, "test.xml")
        generate_xml(jinja_context, path)
        tree = etree.parse(path)
        assert tree.getroot().tag == "IndustrialCatalogue"

    def test_xml_has_product_section(self, jinja_context, tmp_dir):
        """Le XML contient la section Product."""
        from lxml import etree
        path = os.path.join(tmp_dir, "test.xml")
        generate_xml(jinja_context, path)
        tree = etree.parse(path)
        products = tree.findall("Product")
        assert len(products) == 1


# =============================================================================
# Tests des utilitaires (sha256, manifeste)
# =============================================================================

class TestUtils:

    def test_sha256_produces_hash(self, tmp_dir):
        """_sha256 retourne une chaîne hexadécimale de 64 caractères."""
        test_file = os.path.join(tmp_dir, "test.txt")
        with open(test_file, "w") as f:
            f.write("test content")
        h = _sha256(test_file)
        assert len(h) == 64
        assert all(c in "0123456789abcdef" for c in h)

    def test_manifest_structure(self, tmp_dir):
        """Le manifeste a la bonne structure."""
        test_file = os.path.join(tmp_dir, "test.txt")
        with open(test_file, "w") as f:
            f.write("test")
        manifest = _create_manifest({"test.txt": test_file}, "2026-04-12")
        assert "pipeline" in manifest
        assert "files" in manifest
        assert "test.txt" in manifest["files"]
        assert "sha256" in manifest["files"]["test.txt"]


# =============================================================================
# Tests d'intégration — pipeline complet
# =============================================================================

class TestIntegration:

    def test_full_pipeline_runs(self, tmp_dir):
        """Le pipeline complet s'exécute sans erreur."""
        result = run_agent9(FAKE_PIPELINE_DATA, output_dir=tmp_dir)
        assert result["status"] == "success"

    def test_zip_created(self, tmp_dir):
        """Le fichier ZIP est créé."""
        result = run_agent9(FAKE_PIPELINE_DATA, output_dir=tmp_dir)
        assert os.path.exists(result["zip_path"])

    def test_zip_contains_expected_files(self, tmp_dir):
        """Le ZIP contient au moins HTML, Excel, JSON, XML."""
        result = run_agent9(FAKE_PIPELINE_DATA, output_dir=tmp_dir)
        with zipfile.ZipFile(result["zip_path"], "r") as zf:
            names = zf.namelist()
        assert any(".html" in n for n in names)
        assert any(".xlsx" in n for n in names)
        assert any(".json" in n for n in names)
        assert any(".xml" in n for n in names)

    def test_zip_contains_manifest(self, tmp_dir):
        """Le ZIP contient le manifeste."""
        result = run_agent9(FAKE_PIPELINE_DATA, output_dir=tmp_dir)
        with zipfile.ZipFile(result["zip_path"], "r") as zf:
            names = zf.namelist()
        assert "manifeste.json" in names

    def test_invalid_data_raises(self, tmp_dir):
        """Des données invalides lèvent une ValueError."""
        bad_data = {"agent1": {}, "agent2": {}, "agent3": {}, "agent4": {},
                    "agent5": {}, "agent6": {}, "agent7": {}, "agent8": {}}
        with pytest.raises((ValueError, Exception)):
            run_agent9(bad_data, output_dir=tmp_dir)